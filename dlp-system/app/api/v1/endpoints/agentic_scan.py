"""Agentic scan API endpoints."""

from __future__ import annotations

import io
import json
import logging
import re
import threading
from typing import Any

from fastapi import (
    APIRouter,
    BackgroundTasks,
    Depends,
    Header,
    HTTPException,
)
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.core.dependencies import get_current_user
from app.db.session import SessionLocal, get_db
from app.models.enterprise import AgenticScanStatus
from app.models.user import User
from app.repositories.enterprise_repo import EnterpriseRepository
from app.schemas.enterprise import (
    AgenticScanResponse,
    ApprovalRequest,
    ScanGoalRequest,
)
from app.services.agentic_scan_service import AgenticScanService


logger = logging.getLogger(__name__)
scan_logger = logging.getLogger("sentineldlp.scan")

router = APIRouter(
    prefix="/agentic-scans",
    tags=["agentic-scans"],
)

_scan_execution_lock = threading.Lock()
_launch_registry_lock = threading.Lock()
_launch_scan_ids: dict[tuple[int, str], int] = {}

_MAX_LAUNCH_REGISTRY_ITEMS = 1000

_TERMINAL_STATUSES = {
    AgenticScanStatus.COMPLETED,
    AgenticScanStatus.PARTIALLY_COMPLETED,
    AgenticScanStatus.AWAITING_REVIEW,
    AgenticScanStatus.FAILED,
    AgenticScanStatus.CANCELLED,
}

_HEADER_FILL = PatternFill(
    "solid",
    fgColor="0B7285",
)

_HEADER_FONT = Font(
    color="FFFFFF",
    bold=True,
)


def build_scan_response(scan) -> AgenticScanResponse:
    """Convert a persisted scan model into the public API response."""
    status = (
        scan.status.value
        if hasattr(scan.status, "value")
        else str(scan.status)
    )

    workflow_state = scan.workflow_state or {}

    return AgenticScanResponse(
        scan_id=scan.id,
        status=status,
        plan=scan.plan or {},
        summary=scan.summary or {},
        approval_required=scan.approval_required,
        error=scan.error_message,
        runtime=workflow_state.get(
            "runtime",
            "langgraph",
        ),
        thread_id=workflow_state.get("thread_id"),
    )


def get_registered_scan(
    db: Session,
    user_id: int,
    launch_id: str | None,
):
    """Return an existing scan when the browser repeats the same launch."""
    if not launch_id:
        return None

    scan_id = _launch_scan_ids.get(
        (user_id, launch_id),
    )

    if scan_id is None:
        return None

    return EnterpriseRepository(db).get_scan(
        scan_id,
        user_id,
    )


def register_scan_launch(
    user_id: int,
    launch_id: str | None,
    scan_id: int,
) -> None:
    """Keep a bounded server-side idempotency registry."""
    if not launch_id:
        return

    _launch_scan_ids[(user_id, launch_id)] = scan_id

    while len(_launch_scan_ids) > _MAX_LAUNCH_REGISTRY_ITEMS:
        oldest_key = next(iter(_launch_scan_ids))
        _launch_scan_ids.pop(oldest_key, None)


def mark_scan_failed(
    db: Session,
    scan_id: int,
    error: Exception,
) -> None:
    """Persist an unexpected background failure for UI polling."""
    repository = EnterpriseRepository(db)
    scan = repository.get_scan(scan_id)

    if scan is None:
        return

    workflow_state: dict[str, Any] = dict(
        scan.workflow_state or {},
    )

    errors = list(
        workflow_state.get("errors", []),
    )

    errors.append(
        {
            "stage": "background_execution",
            "error": type(error).__name__,
            "message": str(error),
        }
    )

    workflow_state.update(
        {
            "status": AgenticScanStatus.FAILED.value,
            "errors": errors,
            "last_error": errors[-1],
        }
    )

    scan_stats = workflow_state.get("scan_stats") or {}
    scan_logger.error(
        "SCAN_FAILED scan_id=%s status=%s findings=%s files=%s tables=%s "
        "rows=%s error_type=%s message=%s",
        scan_id,
        AgenticScanStatus.FAILED.value,
        len(workflow_state.get("findings") or []),
        scan_stats.get("files_processed", scan_stats.get("files_scanned", 0)),
        scan_stats.get("tables_processed", scan_stats.get("tables_scanned", 0)),
        scan_stats.get("rows_scanned", 0),
        type(error).__name__,
        str(error),
    )

    repository.save_workflow(
        scan,
        workflow_state,
    )


def mark_scan_running(
    db: Session,
    scan_id: int,
) -> None:
    """Persist that a queued scan has started execution."""
    scan = EnterpriseRepository(db).get_scan(scan_id)

    if scan is None:
        return

    workflow_state = dict(
        scan.workflow_state or {},
    )

    workflow_state["status"] = (
        AgenticScanStatus.SCANNING.value
    )

    scan.status = AgenticScanStatus.SCANNING
    scan.workflow_state = workflow_state

    db.commit()


def safe_excel_value(value: Any) -> Any:
    """Convert values and prevent spreadsheet formula injection."""
    if value is None:
        return ""

    if isinstance(value, (int, float, bool)):
        return value

    if isinstance(value, (dict, list, tuple)):
        value = json.dumps(
            value,
            ensure_ascii=False,
            default=str,
        )

    text = str(value)

    if text.startswith(("=", "+", "-", "@")):
        return f"'{text}"

    return text


def format_worksheet(worksheet) -> None:
    """Apply readable formatting to an exported worksheet."""
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for cell in worksheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for column_cells in worksheet.columns:
        column_letter = get_column_letter(
            column_cells[0].column,
        )

        max_length = max(
            len(str(cell.value))
            if cell.value is not None
            else 0
            for cell in column_cells
        )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(max_length + 2, 12),
            60,
        )


def build_scan_workbook(scan) -> Workbook:
    """Build a multi-sheet XLSX report for one scan."""
    workflow_state = scan.workflow_state or {}

    findings = (
        workflow_state.get("verified_findings")
        or workflow_state.get("findings")
        or []
    )

    errors = workflow_state.get("errors") or []

    status = (
        scan.status.value
        if hasattr(scan.status, "value")
        else str(scan.status)
    )

    workbook = Workbook()

    summary_sheet = workbook.active
    summary_sheet.title = "Scan Summary"
    summary_sheet.append(["Field", "Value"])

    summary_rows = [
        ("Scan ID", scan.id),
        (
            "Asset",
            scan.asset.name if scan.asset else "",
        ),
        (
            "Platform",
            scan.asset.platform if scan.asset else "",
        ),
        ("Status", status),
        ("Goal", scan.goal),
        ("Total findings", len(findings)),
        (
            "Approval required",
            scan.approval_required,
        ),
        ("Error", scan.error_message),
        (
            "Risk summary",
            workflow_state.get("risk_summary", {}),
        ),
        (
            "Scan statistics",
            workflow_state.get("scan_stats", {}),
        ),
        (
            "Policy decision",
            workflow_state.get("policy_decision", {}),
        ),
        (
            "Started at",
            getattr(scan, "created_at", None),
        ),
        (
            "Last updated",
            getattr(scan, "updated_at", None),
        ),
    ]

    for field, value in summary_rows:
        summary_sheet.append(
            [
                field,
                safe_excel_value(value),
            ]
        )

    format_worksheet(summary_sheet)

    findings_sheet = workbook.create_sheet(
        "Findings",
    )

    findings_sheet.append(
        [
            "Entity Type",
            "Masked Value",
            "Confidence",
            "Source",
            "Risk Level",
            "Path",
            "SHA256",
            "Location",
        ]
    )

    for finding in findings:
        location = finding.get("location") or {}

        findings_sheet.append(
            [
                safe_excel_value(
                    finding.get("entity_type"),
                ),
                safe_excel_value(
                    finding.get("masked_value"),
                ),
                safe_excel_value(
                    finding.get("confidence"),
                ),
                safe_excel_value(
                    finding.get("source"),
                ),
                safe_excel_value(
                    finding.get("risk_level"),
                ),
                safe_excel_value(
                    location.get("path"),
                ),
                safe_excel_value(
                    location.get("sha256"),
                ),
                safe_excel_value(location),
            ]
        )

    format_worksheet(findings_sheet)

    errors_sheet = workbook.create_sheet(
        "Errors",
    )

    errors_sheet.append(
        [
            "Stage",
            "Error",
            "Message",
            "Details",
        ]
    )

    for error in errors:
        errors_sheet.append(
            [
                safe_excel_value(
                    error.get("stage"),
                ),
                safe_excel_value(
                    error.get("error"),
                ),
                safe_excel_value(
                    error.get("message"),
                ),
                safe_excel_value(error),
            ]
        )

    format_worksheet(errors_sheet)

    return workbook


def execute_scan_in_background(
    scan_id: int,
) -> None:
    """Run a scan after returning its ID to the frontend."""
    with _scan_execution_lock:
        db = SessionLocal()

        try:
            mark_scan_running(
                db,
                scan_id,
            )

            AgenticScanService(db).execute(
                scan_id,
            )
        except Exception as error:
            logger.exception(
                "Background scan %s failed",
                scan_id,
            )

            try:
                mark_scan_failed(
                    db,
                    scan_id,
                    error,
                )
            except Exception:
                logger.exception(
                    "Unable to persist failure for scan %s",
                    scan_id,
                )
        finally:
            db.close()


@router.post(
    "",
    response_model=AgenticScanResponse,
    status_code=201,
)
def start_scan(
    payload: ScanGoalRequest,
    background_tasks: BackgroundTasks,
    launch_id: str | None = Header(
        default=None,
        alias="X-Scan-Launch-Id",
    ),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Create and execute a scan."""
    service = AgenticScanService(db)

    try:
        if launch_id and len(launch_id) > 128:
            raise HTTPException(
                status_code=400,
                detail="Invalid scan launch ID",
            )

        with _launch_registry_lock:
            scan = get_registered_scan(
                db,
                current_user.id,
                launch_id,
            )

            if scan is None:
                scan = service.create(
                    current_user.id,
                    payload,
                )

                register_scan_launch(
                    current_user.id,
                    launch_id,
                    scan.id,
                )
            else:
                return build_scan_response(scan)

        if payload.execute_async:
            background_tasks.add_task(
                execute_scan_in_background,
                scan.id,
            )
            scan_logger.info(
                "SCAN_QUEUED scan_id=%s user_id=%s async=true",
                scan.id,
                current_user.id,
            )
        else:
            scan = service.execute(
                scan.id,
                current_user.id,
            )

        return build_scan_response(scan)

    except ValueError as error:
        raise HTTPException(
            status_code=400,
            detail=str(error),
        ) from error


@router.get("/{scan_id}")
def get_scan(
    scan_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Return the latest scan state."""
    scan = EnterpriseRepository(db).get_scan(
        scan_id,
        current_user.id,
    )

    if scan is None:
        raise HTTPException(
            status_code=404,
            detail="Scan not found",
        )

    return {
        **build_scan_response(scan).model_dump(),
        "workflow_state": scan.workflow_state,
        "agent_runs": [
            {
                "agent": run.agent_name,
                "action": run.action,
                "status": (
                    run.status.value
                    if hasattr(run.status, "value")
                    else str(run.status)
                ),
                "summary": run.output_summary,
            }
            for run in scan.runs
        ],
    }


@router.get("/{scan_id}/export")
def export_scan_report(
    scan_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Download a scan report as XLSX."""
    scan = EnterpriseRepository(db).get_scan(
        scan_id,
        current_user.id,
    )

    if scan is None:
        raise HTTPException(
            status_code=404,
            detail="Scan not found",
        )

    if scan.status not in _TERMINAL_STATUSES:
        raise HTTPException(
            status_code=409,
            detail=(
                "The report will be available "
                "after the scan finishes."
            ),
        )

    workbook = build_scan_workbook(scan)

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    asset_name = (
        scan.asset.name
        if scan.asset
        else "asset"
    )

    safe_asset_name = re.sub(
        r"[^A-Za-z0-9_-]+",
        "-",
        asset_name,
    ).strip("-")

    filename = (
        f"dlp-scan-{scan.id}-"
        f"{safe_asset_name or 'asset'}.xlsx"
    )

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; filename="{filename}"'
            )
        },
    )


@router.post(
    "/{scan_id}/review",
    response_model=AgenticScanResponse,
)
def review_scan(
    scan_id: int,
    payload: ApprovalRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Resume a paused workflow with an analyst decision."""
    try:
        scan = AgenticScanService(db).approve(
            scan_id,
            current_user.id,
            payload.approved,
            payload.note,
        )

        return build_scan_response(scan)

    except ValueError as error:
        raise HTTPException(
            status_code=400,
            detail=str(error),
        ) from error 
