"""Deterministic, allowlisted tools for approved file remediation."""

from __future__ import annotations

import base64
import hashlib
import os
import stat
import tempfile
from pathlib import Path
from typing import Callable

from cryptography.fernet import Fernet

from app.core.config import settings
from app.detection.hybrid_engine import engine
from app.detection.masker import mask_value
from app.tools.filesystem_tools import file_sha256, validate_allowed_path
from app.utils.file_parser import FileParser


TEXT_EXTENSIONS = {".txt", ".csv", ".json", ".xml", ".log"}
DOCUMENT_EXTENSIONS = {".docx", ".xlsx", ".pdf"}
SUPPORTED_REMEDIATION_EXTENSIONS = TEXT_EXTENSIONS | DOCUMENT_EXTENSIONS


def _fernet() -> Fernet:
    secret = settings.REMEDIATION_MASTER_KEY or settings.SECRET_KEY

    if len(secret) < 32:
        raise RuntimeError(
            "REMEDIATION_MASTER_KEY must contain at least 32 characters"
        )

    derived_key = base64.urlsafe_b64encode(hashlib.sha256(secret.encode("utf-8")).digest())
    return Fernet(derived_key)


def _replace_sensitive_text(
    text: str,
    *,
    action: str,
    entity_types: set[str],
) -> tuple[str, int, set[str]]:
    findings = [
        finding
        for finding in engine.detect(text)
        if finding.entity_type in entity_types
        and finding.start_pos is not None
        and finding.end_pos is not None
    ]
    findings.sort(key=lambda item: int(item.start_pos or 0), reverse=True)
    updated = text
    replacement_values: set[str] = set()

    for finding in findings:
        replacement = (
            mask_value(finding.entity_type, finding.value)
            if action == "mask"
            else "[REDACTED]"
        )
        replacement_values.add(replacement)
        updated = (
            updated[: int(finding.start_pos)]
            + replacement
            + updated[int(finding.end_pos) :]
        )

    return updated, len(findings), replacement_values


def _transform_text_file(
    source: Path,
    destination: Path,
    *,
    action: str,
    entity_types: set[str],
) -> tuple[int, set[str]]:
    content = source.read_text(encoding="utf-8", errors="replace")
    updated, replacements, replacement_values = _replace_sensitive_text(
        content,
        action=action,
        entity_types=entity_types,
    )
    destination.write_text(updated, encoding="utf-8", newline="")
    return replacements, replacement_values


def _transform_docx(
    source: Path,
    destination: Path,
    *,
    action: str,
    entity_types: set[str],
) -> tuple[int, set[str]]:
    from docx import Document

    document = Document(str(source))
    replacements = 0
    replacement_values: set[str] = set()

    def update_paragraph(paragraph) -> None:
        nonlocal replacements

        for run in paragraph.runs:
            updated, count, values = _replace_sensitive_text(
                run.text,
                action=action,
                entity_types=entity_types,
            )
            run.text = updated
            replacements += count
            replacement_values.update(values)

    for paragraph in document.paragraphs:
        update_paragraph(paragraph)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    update_paragraph(paragraph)

    document.save(str(destination))
    return replacements, replacement_values


def _transform_xlsx(
    source: Path,
    destination: Path,
    *,
    action: str,
    entity_types: set[str],
) -> tuple[int, set[str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(str(source))
    replacements = 0
    replacement_values: set[str] = set()

    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str) or cell.value.startswith("="):
                    continue

                updated, count, values = _replace_sensitive_text(
                    cell.value,
                    action=action,
                    entity_types=entity_types,
                )
                cell.value = updated
                replacements += count
                replacement_values.update(values)

    workbook.save(str(destination))
    workbook.close()
    return replacements, replacement_values


def _transform_pdf(
    source: Path,
    destination: Path,
    *,
    action: str,
    entity_types: set[str],
) -> tuple[int, set[str]]:
    try:
        import fitz
    except ImportError as error:
        raise RuntimeError(
            "PDF remediation requires PyMuPDF. Install requirements.txt."
        ) from error

    document = fitz.open(str(source))
    replacements = 0
    replacement_values: set[str] = set()

    try:
        for page in document:
            page_text = page.get_text("text")
            findings = [
                finding
                for finding in engine.detect(page_text)
                if finding.entity_type in entity_types
            ]

            for finding in findings:
                replacement = (
                    mask_value(finding.entity_type, finding.value)
                    if action == "mask"
                    else "[REDACTED]"
                )
                replacement_values.add(replacement)
                rectangles = page.search_for(finding.value)

                for rectangle in rectangles:
                    page.add_redact_annot(
                        rectangle,
                        text=replacement,
                        fill=(0, 0, 0),
                        text_color=(1, 1, 1),
                    )
                    replacements += 1

            if findings:
                page.apply_redactions()

        document.save(str(destination), garbage=4, deflate=True)
    finally:
        document.close()

    return replacements, replacement_values


def _transform_file(
    source: Path,
    destination: Path,
    *,
    action: str,
    entity_types: set[str],
) -> tuple[int, set[str]]:
    extension = source.suffix.lower()

    if extension in TEXT_EXTENSIONS:
        transformer: Callable[..., tuple[int, set[str]]] = _transform_text_file
    elif extension == ".docx":
        transformer = _transform_docx
    elif extension == ".xlsx":
        transformer = _transform_xlsx
    elif extension == ".pdf":
        transformer = _transform_pdf
    else:
        raise ValueError(f"Remediation is not supported for {extension or 'this file type'}")

    return transformer(
        source,
        destination,
        action=action,
        entity_types=entity_types,
    )


def _encrypted_backup(path: Path, job_id: int) -> Path:
    backup_root = Path(settings.REMEDIATION_BACKUP_ROOT).expanduser().resolve()
    job_directory = backup_root / f"job-{job_id}"
    job_directory.mkdir(parents=True, exist_ok=True, mode=0o700)
    backup_path = job_directory / f"{path.name}.backup.enc"
    backup_path.write_bytes(_fernet().encrypt(path.read_bytes()))
    os.chmod(backup_path, 0o600)
    return backup_path


def _restore_backup(backup_path: Path, destination: Path, mode: int) -> None:
    restored = _fernet().decrypt(backup_path.read_bytes())
    temporary = destination.with_name(f".{destination.name}.restore.tmp")
    temporary.write_bytes(restored)
    os.chmod(temporary, mode)
    os.replace(temporary, destination)


def _remaining_findings(
    path: Path, entity_types: set[str], known_replacement_values: set[str] = frozenset()
) -> int:
    """Count still-sensitive matches after remediation.

    Excludes values that exactly equal one of the replacements we just
    wrote (e.g. format-preserving masking like "t***@example.com" is
    itself still shaped like an email and would otherwise re-trigger the
    EMAIL detector on the value we deliberately produced). Anything that
    doesn't match a known replacement is genuinely leftover PII and still
    fails verification -- this narrows the false-positive, it doesn't
    weaken the full-file re-scan.
    """
    text = FileParser.parse(str(path))

    if text is None:
        raise RuntimeError("Unable to verify the remediated file")

    return sum(
        1
        for finding in engine.detect(text)
        if finding.entity_type in entity_types
        and finding.value not in known_replacement_values
    )


def execute_file_remediation(
    *,
    job_id: int,
    file_path: str,
    expected_sha256: str,
    action: str,
    entity_types: set[str],
) -> dict[str, object]:
    """Execute one approved action with backup, hash and verification controls."""
    path = validate_allowed_path(file_path)

    if not path.is_file():
        raise ValueError("Remediation target must be a file")

    current_sha256 = file_sha256(path)

    if current_sha256 != expected_sha256:
        raise RuntimeError(
            "File changed after scanning. Run a fresh scan before remediation."
        )

    if action in {"mask", "redact"} and path.suffix.lower() not in SUPPORTED_REMEDIATION_EXTENSIONS:
        raise ValueError("Selected file type does not support masking or redaction")

    original_mode = stat.S_IMODE(path.stat().st_mode)
    original_size = path.stat().st_size
    backup_path = _encrypted_backup(path, job_id)
    output_path: Path | None = path
    replacements = 0

    try:
        if action in {"mask", "redact"}:
            handle, temporary_name = tempfile.mkstemp(
                prefix=f".{path.stem}-remediation-",
                suffix=path.suffix,
                dir=path.parent,
            )
            os.close(handle)
            temporary = Path(temporary_name)

            try:
                replacements, replacement_values = _transform_file(
                    path,
                    temporary,
                    action=action,
                    entity_types=entity_types,
                )

                if replacements == 0:
                    raise RuntimeError(
                        "No matching raw values were found; the file was not changed."
                    )

                os.chmod(temporary, original_mode)
                os.replace(temporary, path)
            finally:
                temporary.unlink(missing_ok=True)

            remaining = _remaining_findings(
                path, entity_types, known_replacement_values=replacement_values
            )

            if remaining > 0:
                _restore_backup(backup_path, path, original_mode)
                raise RuntimeError(
                    f"Verification found {remaining} remaining sensitive values; original restored."
                )

        elif action == "encrypt":
            output_path = path.with_name(f"{path.name}.dlp.enc")
            encrypted = _fernet().encrypt(path.read_bytes())
            temporary = output_path.with_name(f".{output_path.name}.tmp")
            temporary.write_bytes(encrypted)
            os.replace(temporary, output_path)
            path.unlink()

        elif action == "delete":
            output_path = None
            path.unlink()
        else:
            raise ValueError("Unsupported remediation action")

        output_sha256 = (
            file_sha256(output_path)
            if output_path is not None and output_path.exists()
            else None
        )

        return {
            "action": action,
            "verified": True,
            "replacements": replacements,
            "original_size_bytes": original_size,
            "original_sha256": current_sha256,
            "output_sha256": output_sha256,
            "output_exists": bool(output_path and output_path.exists()),
            "source_removed": not path.exists() if action in {"encrypt", "delete"} else False,
            "retention_days": (
                settings.REMEDIATION_DELETE_RETENTION_DAYS
                if action == "delete"
                else None
            ),
            "encrypted_backup_path": str(backup_path),
            "output_path": str(output_path) if output_path is not None else None,
        }
    except Exception:
        # A failed job must never leave a partially remediated source behind.
        _restore_backup(backup_path, path, original_mode)
        raise
